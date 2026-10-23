# ============================================================
# STREAMLIT RAG CHATBOT
# WORD / PDF / HTML FILES + MILVUS + OLLAMA (llama3.2:3b)
# HYBRID SEARCH: VECTOR SEARCH + KEYWORD SEARCH + EMAIL BOOST
# ============================================================
#
# Required packages:
#
# pip install streamlit python-docx pypdf pymilvus
# pip install sentence-transformers ollama
# pip install beautifulsoup4          <-- NEW (for .html files)
#
# Make sure Milvus is running on 127.0.0.1:19530
# Make sure Ollama has the model pulled:
#   ollama pull llama3.2:3b
#
# Put your .html / .docx / .pdf files in a folder named
# "documents" next to this script, then run:
#   streamlit run streamlit_rag_app.py
#
# ============================================================

import os
import re

import streamlit as st
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from pymilvus import (
    connections,
    Collection,
    FieldSchema,
    CollectionSchema,
    DataType,
    utility
)

from sentence_transformers import SentenceTransformer

import ollama


# ============================================================
# CONFIGURATION
# ============================================================

COLLECTION_NAME = "StreamlitRAG"

CHUNK_SIZE = 150
CHUNK_OVERLAP = 30

EMBEDDING_MODEL = "all-MiniLM-L6-v2"
OLLAMA_MODEL = "llama3.2:3b"

MILVUS_HOST = os.environ.get("MILVUS_HOST", "127.0.0.1")
MILVUS_PORT = os.environ.get("MILVUS_PORT", "19530")

MAX_FILES = 10

DEFAULT_DOCUMENT_FOLDER = "documents"

# File types the app will pick up from the folder
SUPPORTED_EXTENSIONS = (".docx", ".pdf", ".html", ".htm")

EMAIL_MATCH_BONUS = 5
NAME_TOKEN_BONUS = 1

EMAIL_PATTERN = r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"


# ============================================================
# CACHED RESOURCES
# ============================================================

@st.cache_resource(show_spinner=False)
def get_embedding_model():
    return SentenceTransformer(EMBEDDING_MODEL)


@st.cache_resource(show_spinner=False)
def get_milvus_connection():
    connections.connect(
        alias="default",
        host=MILVUS_HOST,
        port=MILVUS_PORT
    )
    return True


# ============================================================
# TEXT EXTRACTION
# ============================================================

def extract_docx_text(file_obj):

    document = Document(file_obj)

    full_text = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            full_text.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if row_text:
                full_text.append(" | ".join(row_text))

    complete_text = "\n".join(full_text)
    complete_text = re.sub(r"\s+", " ", complete_text).strip()

    return complete_text


def extract_pdf_text(file_obj):

    reader = PdfReader(file_obj)

    full_text = []

    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            full_text.append(text)

    complete_text = "\n".join(full_text)
    complete_text = re.sub(r"\s+", " ", complete_text).strip()

    return complete_text


def extract_html_text(source):
    """
    Extract readable text from an HTML file.

    `source` can be a file path (str) or a file-like object
    (e.g. a Streamlit UploadedFile).

    - Strips <script>, <style>, <head>, <noscript> content
    - Keeps table rows together as "cell | cell" lines, so
      e.g. "Cloud Platform | Azure, Oracle" stays in one chunk
    - Keeps mailto: links, so email addresses hidden in an
      <a href="mailto:..."> tag are not lost
    """

    if isinstance(source, str):
        with open(source, "rb") as f:
            raw = f.read()
    else:
        raw = source.read()

    # Decode safely (utf-8 first, fall back to latin-1)
    try:
        html = raw.decode("utf-8")
    except UnicodeDecodeError:
        html = raw.decode("latin-1", errors="replace")

    soup = BeautifulSoup(html, "html.parser")

    # Remove non-content tags
    for tag in soup(["script", "style", "head", "noscript"]):
        tag.decompose()

    # Keep mailto links visible in the text
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if href.lower().startswith("mailto:"):
            email = href[7:].split("?")[0].strip()
            if email and email not in a.get_text():
                a.append(f" ({email})")

    # Turn every table row into one "cell | cell | cell" line
    for table in soup.find_all("table"):
        lines = []
        for row in table.find_all("tr"):
            cells = [
                c.get_text(" ", strip=True)
                for c in row.find_all(["th", "td"])
            ]
            cells = [c for c in cells if c]
            if cells:
                lines.append(" | ".join(cells))
        table.replace_with(soup.new_string("\n" + "\n".join(lines) + "\n"))

    complete_text = soup.get_text("\n", strip=True)
    complete_text = re.sub(r"\s+", " ", complete_text).strip()

    return complete_text


def extract_text(source):
    """
    Extract text from either a Streamlit UploadedFile or a
    plain file path string. Supports .docx, .pdf, .html, .htm
    """

    if isinstance(source, str):
        name = source.lower()
    else:
        name = source.name.lower()

    if name.endswith(".docx"):
        return extract_docx_text(source)

    if name.endswith(".pdf"):
        return extract_pdf_text(source)

    if name.endswith((".html", ".htm")):
        return extract_html_text(source)

    return ""


def display_name(source):
    if isinstance(source, str):
        return os.path.basename(source)
    return source.name


# ============================================================
# LOCAL FOLDER DOCUMENT FINDING
# ============================================================

def find_document_files(folder_path):
    """
    Scan a local folder for .docx, .pdf, .html and .htm files.
    Returns a list of full file paths.
    """

    if not folder_path:
        return []

    if not os.path.exists(folder_path):
        return []

    if not os.path.isdir(folder_path):
        return []

    found_files = []

    for file in sorted(os.listdir(folder_path)):

        if file.lower().endswith(SUPPORTED_EXTENSIONS):
            found_files.append(os.path.join(folder_path, file))

    return found_files


# ============================================================
# CHUNKING
# ============================================================

def chunk_document(text, file_name, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):

    words = text.split()

    chunks = []
    start = 0

    while start < len(words):

        end = start + chunk_size
        chunk_words = words[start:end]
        chunk_text = " ".join(chunk_words)

        if chunk_text.strip():
            chunks.append({
                "document": file_name,
                "text": chunk_text
            })

        start += (chunk_size - overlap)

    return chunks


# ============================================================
# EMAIL / NAME TOKEN HELPERS
# ============================================================

def normalize_text(text):
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def get_name_tokens(text):

    q = normalize_text(text)

    stop_words = {
        "what", "is", "the", "email", "emailid",
        "emailaddress", "address", "id", "of", "for",
        "please", "tell", "me", "give", "find", "show",
        "mail", "mailid", "mailaddress"
    }

    words = re.findall(r"[a-zA-Z0-9._+-]+", q)

    return [w for w in words if w not in stop_words]


def score_email_matches(chunk_text, name_tokens):

    emails_found = re.findall(EMAIL_PATTERN, chunk_text)

    if not emails_found:
        return 0

    normalized_chunk = normalize_text(chunk_text)

    matched_tokens = sum(
        1 for token in name_tokens
        if token in normalized_chunk
    )

    return EMAIL_MATCH_BONUS + (matched_tokens * NAME_TOKEN_BONUS)


# ============================================================
# KEYWORD SEARCH
# ============================================================

def keyword_search(query, chunks, top_k=10):

    query_words = set(
        re.findall(r"\b[a-zA-Z0-9@._+-]+\b", query.lower())
    )

    results = []

    for chunk in chunks:

        text = chunk["text"].lower()

        matched_words = [w for w in query_words if w in text]
        score = len(matched_words)

        if score > 0:
            results.append({
                "document": chunk["document"],
                "text": chunk["text"],
                "keyword_score": score
            })

    results.sort(key=lambda x: x["keyword_score"], reverse=True)

    return results[:top_k]


# ============================================================
# MILVUS INDEXING
# ============================================================

def build_milvus_collection(all_chunks, embedding_model):

    get_milvus_connection()

    if utility.has_collection(COLLECTION_NAME):
        Collection(COLLECTION_NAME).drop()

    fields = [
        FieldSchema(name="id", dtype=DataType.INT64, is_primary=True, auto_id=False),
        FieldSchema(name="document", dtype=DataType.VARCHAR, max_length=500),
        FieldSchema(name="text", dtype=DataType.VARCHAR, max_length=10000),
        FieldSchema(name="embedding", dtype=DataType.FLOAT_VECTOR, dim=384),
    ]

    schema = CollectionSchema(fields, description="Streamlit RAG Documents")

    collection = Collection(name=COLLECTION_NAME, schema=schema)

    chunk_texts = [c["text"] for c in all_chunks]
    document_names = [c["document"] for c in all_chunks]
    ids = list(range(1, len(chunk_texts) + 1))

    embeddings = embedding_model.encode(
        chunk_texts,
        convert_to_numpy=True,
        show_progress_bar=False
    )

    collection.insert([ids, document_names, chunk_texts, embeddings])
    collection.flush()

    collection.create_index(
        field_name="embedding",
        index_params={
            "metric_type": "COSINE",
            "index_type": "IVF_FLAT",
            "params": {"nlist": 128}
        }
    )

    collection.load()

    return collection


# ============================================================
# RAG QUERY (hybrid vector + keyword + email bonus)
# ============================================================

def run_rag_query(query, all_chunks, collection, embedding_model):

    is_email_question = any(
        keyword in normalize_text(query)
        for keyword in [
            "email", "email id", "email address",
            "mail id", "mail address"
        ]
    )

    query_name_tokens = get_name_tokens(query) if is_email_question else []

    keyword_results = keyword_search(query, all_chunks, top_k=10)

    query_vector = embedding_model.encode(query).tolist()

    search_results = collection.search(
        data=[query_vector],
        anns_field="embedding",
        param={"metric_type": "COSINE", "params": {"nprobe": 10}},
        limit=10,
        output_fields=["document", "text"]
    )

    combined_results = []

    for hit in search_results[0]:
        combined_results.append({
            "document": hit.entity.get("document"),
            "text": hit.entity.get("text"),
            "vector_score": float(hit.distance),
            "keyword_score": 0,
            "email_score": 0
        })

    for kr in keyword_results:

        found = False

        for item in combined_results:
            if item["document"] == kr["document"] and item["text"] == kr["text"]:
                item["keyword_score"] = kr["keyword_score"]
                found = True
                break

        if not found:
            combined_results.append({
                "document": kr["document"],
                "text": kr["text"],
                "vector_score": 0,
                "keyword_score": kr["keyword_score"],
                "email_score": 0
            })

    if is_email_question:
        for item in combined_results:
            item["email_score"] = score_email_matches(item["text"], query_name_tokens)

    combined_results.sort(
        key=lambda x: (x["email_score"], x["keyword_score"], x["vector_score"]),
        reverse=True
    )

    combined_results = combined_results[:10]

    contexts = []

    for item in combined_results:
        contexts.append(
            f"\nDocument: {item['document']}\n\nContent:\n{item['text']}\n"
        )

    context = "\n\n".join(contexts)

    prompt = f"""
You are a document question-answering assistant.

You have access to information retrieved from
multiple Word/PDF/HTML documents.

Your task is to answer the user's question using
ONLY the information provided in the context below.

IMPORTANT RULES:

1. Use ONLY the provided context.
2. Do NOT use your general knowledge.
3. Do NOT invent information.
4. Do NOT guess.
5. If the question asks for an exact value such as:
   - email address
   - phone number
   - URL
   - name
   - date
   - certification number
   - company name

   copy the value EXACTLY as it appears in the provided context.

6. Do not change the spelling, capitalization,
   numbers, or symbols of an exact value.
7. Mention the document name containing the answer.
8. If the answer is not available in the context, respond exactly:

   "The information is not available in the provided documents."

==================================================
RETRIEVED CONTEXT
==================================================

{context}

==================================================
USER QUESTION
==================================================

{query}

==================================================

ANSWER:
"""

    response = ollama.chat(
        model=OLLAMA_MODEL,
        messages=[{"role": "user", "content": prompt}]
    )

    answer = response["message"]["content"]

    return answer, combined_results


# ============================================================
# STREAMLIT UI
# ============================================================

st.set_page_config(page_title="Document RAG Chatbot", page_icon="📄", layout="wide")

st.title("📄 Document RAG Chatbot")
st.caption("Ask questions about your local documents, answered by Llama 3.2 (3B) via Ollama.")

if "processed" not in st.session_state:
    st.session_state.processed = False
if "all_chunks" not in st.session_state:
    st.session_state.all_chunks = []
if "collection" not in st.session_state:
    st.session_state.collection = None
if "messages" not in st.session_state:
    st.session_state.messages = []

folder_files = find_document_files(DEFAULT_DOCUMENT_FOLDER)

if len(folder_files) > MAX_FILES:
    folder_files = folder_files[:MAX_FILES]

should_process = folder_files and not st.session_state.processed

if should_process:

    with st.spinner("Setting things up..."):

        all_chunks = []

        for source in folder_files:

            name = display_name(source)

            try:
                text = extract_text(source)
            except Exception as e:
                st.warning(f"Could not read {name}: {e}")
                continue

            if not text:
                continue

            chunks = chunk_document(text, name)
            all_chunks.extend(chunks)

        if all_chunks:

            embedding_model = get_embedding_model()

            try:
                collection = build_milvus_collection(all_chunks, embedding_model)

                st.session_state.all_chunks = all_chunks
                st.session_state.collection = collection
                st.session_state.processed = True
                st.session_state.messages = []

            except Exception as e:
                st.error(f"Could not connect to or index in Milvus: {e}")


st.header("Chat")

if not st.session_state.processed:

    st.info(
        f"No documents found yet. Add .html, .docx or .pdf files to the "
        f"'{DEFAULT_DOCUMENT_FOLDER}' folder and refresh the page."
    )

else:

    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])
            if msg["role"] == "assistant" and msg.get("sources"):
                with st.expander("Retrieved sources"):
                    for i, s in enumerate(msg["sources"], 1):
                        st.markdown(
                            f"**Source {i} — {s['document']}**  \n"
                            f"vector: {s['vector_score']:.4f} | "
                            f"keyword: {s['keyword_score']} | "
                            f"email bonus: {s['email_score']}"
                        )
                        st.text(s["text"][:500])

    question = st.chat_input("Ask a question about your documents...")

    if question:

        st.session_state.messages.append({"role": "user", "content": question})

        with st.chat_message("user"):
            st.markdown(question)

        with st.chat_message("assistant"):

            with st.spinner("Thinking..."):

                try:
                    embedding_model = get_embedding_model()
                    answer, sources = run_rag_query(
                        question,
                        st.session_state.all_chunks,
                        st.session_state.collection,
                        embedding_model
                    )
                except Exception as e:
                    answer = f"Error while answering: {e}"
                    sources = []

            st.markdown(answer)

            if sources:
                with st.expander("Retrieved sources"):
                    for i, s in enumerate(sources, 1):
                        st.markdown(
                            f"**Source {i} — {s['document']}**  \n"
                            f"vector: {s['vector_score']:.4f} | "
                            f"keyword: {s['keyword_score']} | "
                            f"email bonus: {s['email_score']}"
                        )
                        st.text(s["text"][:500])

        st.session_state.messages.append({
            "role": "assistant",
            "content": answer,
            "sources": sources
        })
